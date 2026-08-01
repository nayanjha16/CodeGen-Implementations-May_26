"""Generate full and short PowerPoint decks for LoRA v1–v3 vs baseline comparison."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMAGES = ROOT / "docs" / "reference" / "images" / "lora-v1-v3"
OUT_FULL = ROOT / "docs" / "reference" / "lora-v1-v3-vs-baseline-comparison.pptx"
OUT_SHORT = ROOT / "docs" / "reference" / "lora-v1-v3-vs-baseline-comparison-short.pptx"

# Palette
NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE_DARK = RGBColor(0x1E, 0x40, 0xAF)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
BLUE_SOFT = RGBColor(0xDB, 0xEA, 0xFE)
BLUE_LIGHT = RGBColor(0x93, 0xC5, 0xFD)
GREEN = RGBColor(0x05, 0x96, 0x69)
GREEN_SOFT = RGBColor(0xD1, 0xFA, 0xE5)
SLATE = RGBColor(0x64, 0x74, 0x8B)
SLATE_LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
TEXT = RGBColor(0x1E, 0x29, 0x3B)
TEXT_MUTED = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xD9, 0x77, 0x06)


def _font(run, *, size=18, bold=False, color=TEXT, name="Segoe UI") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def _accent_bar(slide, height=Inches(0.12)) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE_DARK
    bar.line.fill.background()


def _footer(slide, text: str = "CodeGen LoRA · Spider Gold Validation · 2026") -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(9), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    _font(run, size=9, color=SLATE)


def _header(slide, title: str, subtitle: str = "") -> None:
    _accent_bar(slide)
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(8.9), Inches(0.65))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, size=30, bold=True, color=NAVY)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.95), Inches(8.9), Inches(0.45))
        sp = sub_box.text_frame.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        _font(srun, size=15, color=TEXT_MUTED)


def _rounded_card(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _kpi_card(
    slide,
    left,
    top,
    width,
    height,
    label: str,
    value: str,
    detail: str = "",
    accent: RGBColor = BLUE_DARK,
    fill: RGBColor = SLATE_LIGHT,
) -> None:
    _rounded_card(slide, left, top, width, height, fill, BLUE_SOFT)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    label_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.18), width - Inches(0.35), Inches(0.35))
    lp = label_box.text_frame.paragraphs[0]
    lrun = lp.add_run()
    lrun.text = label.upper()
    _font(lrun, size=11, bold=True, color=SLATE)

    val_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.5), width - Inches(0.35), Inches(0.75))
    vp = val_box.text_frame.paragraphs[0]
    vrun = vp.add_run()
    vrun.text = value
    _font(vrun, size=34, bold=True, color=NAVY)

    if detail:
        det_box = slide.shapes.add_textbox(left + Inches(0.22), top + height - Inches(0.55), width - Inches(0.35), Inches(0.4))
        dp = det_box.text_frame.paragraphs[0]
        drun = dp.add_run()
        drun.text = detail
        _font(drun, size=12, color=TEXT_MUTED)


def _insight_box(slide, text: str, top=Inches(6.55)) -> None:
    box = _rounded_card(slide, Inches(0.55), top, Inches(8.9), Inches(0.55), BLUE_SOFT, BLUE_LIGHT)
    tf = slide.shapes.add_textbox(Inches(0.75), top + Inches(0.1), Inches(8.5), Inches(0.4)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Insight: {text}"
    _font(run, size=13, bold=True, color=BLUE_DARK)


def _bullets(slide, items: list[str], top=Inches(1.45), left=Inches(0.65), width=Inches(8.7), size=19) -> None:
    box = slide.shapes.add_textbox(left, top, width, Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = item
        _font(run, size=size, color=TEXT)


def _chart(slide, image_name: str, top=Inches(1.35), width=Inches(8.9)) -> None:
    path = IMAGES / image_name
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(0.55), top, width=width)


def _simple_table(
    slide,
    headers: list[str],
    rows: list[list[str]],
    top=Inches(1.45),
    highlight_rows: set[int] | None = None,
    col_widths: list[float] | None = None,
) -> None:
    highlight_rows = highlight_rows or set()
    n_rows = len(rows) + 1
    n_cols = len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.55), top, Inches(8.9), Inches(0.52 * n_rows))
    table = shape.table

    if col_widths:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Inches(w)
    else:
        w = 8.9 / n_cols
        for idx in range(n_cols):
            table.columns[idx].width = Inches(w)

    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for para in cell.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                _font(run, size=12, bold=True, color=WHITE)

    for r, row in enumerate(rows, start=1):
        highlight = r in highlight_rows
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if highlight:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE_SOFT
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = SLATE_LIGHT
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                for run in para.runs:
                    _font(run, size=12, bold=highlight and c == 0, color=BLUE_DARK if highlight and c > 0 else TEXT)


# ── Shared slides ─────────────────────────────────────────────────────────────

def _slide_title(prs: Presentation, tagline: str = "") -> None:
    slide = _blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.5), Inches(-1.2), Inches(4.5), Inches(4.5))
    glow.fill.solid()
    glow.fill.fore_color.rgb = BLUE_DARK
    glow.line.fill.background()

    badge = _rounded_card(slide, Inches(0.65), Inches(0.65), Inches(2.2), Inches(0.42), BLUE_DARK)
    btf = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(2), Inches(0.3)).text_frame
    bp = btf.paragraphs[0]
    br = bp.add_run()
    br.text = "CAPSTONE RESULTS"
    _font(br, size=10, bold=True, color=WHITE)

    title = slide.shapes.add_textbox(Inches(0.65), Inches(1.35), Inches(8.5), Inches(1.5))
    tp = title.text_frame.paragraphs[0]
    tr = tp.add_run()
    tr.text = "LoRA Fine-Tuning\nBeats Baseline"
    _font(tr, size=44, bold=True, color=WHITE)

    sub = slide.shapes.add_textbox(Inches(0.65), Inches(3.05), Inches(8.5), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sr = sp.add_run()
    sr.text = tagline or "CodeGen-350M · Spider Gold · 50 examples · 3 tasks"
    _font(sr, size=20, color=BLUE_LIGHT)

    chips = ["Text2SQL", "SQL2NoSQL", "Documentation"]
    for i, chip in enumerate(chips):
        left = Inches(0.65 + i * 2.05)
        _rounded_card(slide, left, Inches(4.1), Inches(1.85), Inches(0.48), BLUE_DARK)
        ctf = slide.shapes.add_textbox(left + Inches(0.15), Inches(4.2), Inches(1.55), Inches(0.3)).text_frame
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = chip
        _font(cr, size=12, bold=True, color=WHITE)

    meta = slide.shapes.add_textbox(Inches(0.65), Inches(5.0), Inches(8.5), Inches(1.2))
    lines = [
        "Model: Salesforce/codegen-350M-multi",
        "Best adapter: LoRA v3 (full TEND ~8k, r=32+FFN, 10 epochs)",
        "Judge: gemma3:4b · Evaluated with live database execution",
    ]
    mtf = meta.text_frame
    for i, line in enumerate(lines):
        mp = mtf.paragraphs[0] if i == 0 else mtf.add_paragraph()
        mr = mp.add_run()
        mr.text = line
        _font(mr, size=13, color=RGBColor(0xCB, 0xD5, 0xE1))


def _slide_problem(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    _header(slide, "The starting point", "Zero-shot CodeGen struggles on real database tasks")
    _kpi_card(slide, Inches(0.55), Inches(1.55), Inches(2.85), Inches(1.65), "Text2SQL", "14%", "queries run correctly", SLATE, SLATE_LIGHT)
    _kpi_card(slide, Inches(3.55), Inches(1.55), Inches(2.85), Inches(1.65), "SQL2NoSQL", "22%", "queries run correctly", SLATE, SLATE_LIGHT)
    _kpi_card(slide, Inches(6.55), Inches(1.55), Inches(2.85), Inches(1.65), "Exact match", "0–4%", "string-perfect output", SLATE, SLATE_LIGHT)
    _bullets(
        slide,
        [
            "Baseline = same model with no LoRA adapter",
            "We tested on 50 curated Spider gold examples",
            "Success = generated query actually runs and returns correct results",
            "Goal: show LoRA training improves all three pipeline stages",
        ],
        top=Inches(3.55),
        size=18,
    )
    _footer(slide)


def _slide_journey(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    _header(slide, "Training journey", "Smoke test → full TEND → wider LoRA (v3)")

    steps = [
        ("Baseline", "No training", "14% / 22%", SLATE_LIGHT, SLATE),
        ("LoRA v1", "50 samples · smoke test", "Pipeline check only", SLATE_LIGHT, AMBER),
        ("LoRA v2", "~8k · r=16 · 5 epochs", "60% / 74%", BLUE_SOFT, BLUE),
        ("LoRA v3", "~8k · r=32+FFN · 10 ep", "66% / 86%", GREEN_SOFT, GREEN),
    ]
    for i, (name, train, result, fill, accent) in enumerate(steps):
        left = Inches(0.55 + i * 2.35)
        _rounded_card(slide, left, Inches(1.55), Inches(2.15), Inches(2.35), fill, accent)
        nb = slide.shapes.add_textbox(left + Inches(0.15), Inches(1.75), Inches(1.85), Inches(0.4))
        nr = nb.text_frame.paragraphs[0].add_run()
        nr.text = name
        _font(nr, size=16, bold=True, color=NAVY)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.25), Inches(1.85), Inches(0.55))
        tr = tb.text_frame.paragraphs[0].add_run()
        tr.text = train
        _font(tr, size=11, color=TEXT_MUTED)

        rb = slide.shapes.add_textbox(left + Inches(0.15), Inches(3.05), Inches(1.85), Inches(0.55))
        rr = rb.text_frame.paragraphs[0].add_run()
        rr.text = result
        _font(rr, size=20, bold=True, color=accent)

        if i < len(steps) - 1:
            arrow = slide.shapes.add_textbox(left + Inches(2.05), Inches(2.45), Inches(0.35), Inches(0.4))
            ar = arrow.text_frame.paragraphs[0].add_run()
            ar.text = "→"
            _font(ar, size=22, bold=True, color=SLATE)

    _insight_box(
        slide,
        "Use v2/v3 for capstone numbers. v1 (n=5) only proves the evaluation pipeline works.",
        top=Inches(4.25),
    )
    _footer(slide)


def _slide_headline_kpis(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    _header(slide, "Bottom line", "LoRA v3 vs baseline on the same 50 examples")

    cards = [
        ("Text2SQL accuracy", "14% → 66%", "+52 points", GREEN),
        ("SQL2NoSQL accuracy", "22% → 86%", "+64 points", GREEN),
        ("Doc quality (judge)", "8.33 → 8.82", "+0.49", BLUE),
    ]
    for i, (label, value, delta, accent) in enumerate(cards):
        left = Inches(0.55 + i * 3.05)
        _kpi_card(slide, left, Inches(1.55), Inches(2.85), Inches(2.0), label, value, delta, accent, SLATE_LIGHT)

    _simple_table(
        slide,
        ["Version", "Training data", "Text2SQL", "SQL2NoSQL", "Doc judge"],
        [
            ["Baseline", "—", "14%", "22%", "8.33"],
            ["LoRA v2", "Full TEND, r=16, 5 ep", "60%", "74%", "8.41"],
            ["LoRA v3 ★", "Full TEND, r=32+FFN, 10 ep", "66%", "86%", "8.82"],
        ],
        top=Inches(3.85),
        highlight_rows={3},
        col_widths=[1.6, 2.0, 1.5, 1.7, 1.5],
    )
    _footer(slide)


def _slide_chart(
    prs: Presentation,
    title: str,
    subtitle: str,
    image: str,
    insight: str,
    chart_top=Inches(1.35),
) -> None:
    slide = _blank_slide(prs)
    _header(slide, title, subtitle)
    _chart(slide, image, top=chart_top, width=Inches(8.85))
    _insight_box(slide, insight, top=Inches(6.35))
    _footer(slide)


def _slide_v3_hero(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    _header(slide, "Recommended model: LoRA v3", "Deploy to Hugging Face Hub + Cloud Run")
    _chart(slide, "07_v3_baseline_vs_lora_summary.png", top=Inches(1.25), width=Inches(6.2))

    _rounded_card(slide, Inches(6.95), Inches(1.35), Inches(2.5), Inches(4.85), GREEN_SOFT, GREEN)
    checks = [
        "Best execution accuracy",
        "Best exact match",
        "Best doc judge score",
        "Trained on full TEND scale",
        "Production-ready adapter",
    ]
    y = Inches(1.65)
    for item in checks:
        box = slide.shapes.add_textbox(Inches(7.15), y, Inches(2.2), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = f"✓  {item}"
        _font(run, size=13, bold=True, color=GREEN)
        y += Inches(0.72)

    _insight_box(slide, "Set manifest.yaml → checkpoint_version: v3", top=Inches(6.35))
    _footer(slide)


def _slide_takeaways(prs: Presentation, short: bool = False) -> None:
    slide = _blank_slide(prs)
    _header(slide, "Key takeaways", "What to say in your capstone presentation")

    items = [
        "LoRA clearly beats zero-shot CodeGen — the improvement is large and consistent.",
        "Biggest wins: SQL2NoSQL (22% → 86%) and exact match (4% → 78%).",
        "Documentation improves modestly; code generation benefits most from fine-tuning.",
        "LoRA capacity and schedule: v2 (r=16, 5 epochs) → v3 (r=32 + FFN, 10 epochs) on the same full TEND data.",
        "Deploy LoRA v3 as the production adapter for the agent and API.",
    ]
    if short:
        items = items[:4]

    _bullets(slide, [f"{i + 1}.  {t}" for i, t in enumerate(items)], top=Inches(1.5), size=20)

    rec = _rounded_card(slide, Inches(0.55), Inches(5.75), Inches(8.9), Inches(0.85), NAVY)
    rtf = slide.shapes.add_textbox(Inches(0.8), Inches(5.95), Inches(8.4), Inches(0.5)).text_frame
    rp = rtf.paragraphs[0]
    rp.alignment = PP_ALIGN.CENTER
    rr = rp.add_run()
    rr.text = "Recommendation → Promote LoRA v3 to production"
    _font(rr, size=18, bold=True, color=WHITE)
    _footer(slide)


# ── Deck builders ─────────────────────────────────────────────────────────────

def build_short_presentation() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _slide_title(prs, "Executive summary · 5-minute version")
    _slide_problem(prs)
    _slide_journey(prs)
    _slide_chart(
        prs,
        "Main result",
        "Execution accuracy improves at every training stage",
        "00_improvement_execution_baseline_v1_v2_v3.png",
        "Text2SQL: 14% → 66% · SQL2NoSQL: 22% → 86% with LoRA v3",
        chart_top=Inches(1.3),
    )
    _slide_v3_hero(prs)
    _slide_takeaways(prs, short=True)

    OUT_SHORT.parent.mkdir(parents=True, exist_ok=True)
    return _save_prs(prs, OUT_SHORT)


def build_full_presentation() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _slide_title(prs)
    _slide_problem(prs)
    _slide_journey(prs)
    _slide_headline_kpis(prs)

    _slide_chart(
        prs,
        "Execution accuracy",
        "Does the generated query run and return correct results?",
        "00_improvement_execution_baseline_v1_v2_v3.png",
        "Baseline 14% Text2SQL → LoRA v3 66%. SQL2NoSQL jumps from 22% to 86%.",
    )
    _slide_chart(
        prs,
        "Exact match",
        "How often is the output identical to the gold answer?",
        "00_improvement_exact_match_baseline_v1_v2_v3.png",
        "Exact match was near zero for baseline; LoRA v3 reaches 48% (SQL) and 78% (NoSQL).",
    )
    _slide_chart(
        prs,
        "Documentation quality",
        "LLM judge scores documentation on a 0–10 scale",
        "00_improvement_doc_judge_baseline_v2_v3.png",
        "Documentation improves steadily but modestly — code tasks gain the most from LoRA.",
    )
    _slide_v3_hero(prs)

    slide = _blank_slide(prs)
    _header(slide, "Baseline vs LoRA — side by side", "Paired evaluation on v2 and v3 (50 examples each)")
    _chart(slide, "01_execution_accuracy_baseline_vs_lora.png", top=Inches(1.25), width=Inches(4.35))
    _chart(slide, "02_exact_match_baseline_vs_lora.png", top=Inches(1.25), width=Inches(4.35))
    pic2 = slide.shapes[-1]
    pic2.left = Inches(5.1)
    _footer(slide)

    slide = _blank_slide(prs)
    _header(slide, "Optional: v1 smoke test", "Pipeline validation only — 5 examples, not for final metrics")
    _chart(slide, "06_v1_smoke_judge_correct_rate.png", top=Inches(1.35), width=Inches(7.5))
    _insight_box(slide, "Mention v1 only to show the eval pipeline works early in the project.", top=Inches(6.35))
    _footer(slide)

    _slide_takeaways(prs, short=False)

    OUT_FULL.parent.mkdir(parents=True, exist_ok=True)
    return _save_prs(prs, OUT_FULL)


def _save_prs(prs: Presentation, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}-new{path.suffix}")
        prs.save(str(alt))
        print(f"Note: {path.name} is open — wrote {alt.name} instead. Close the file and re-run to overwrite.")
        return alt


if __name__ == "__main__":
    full = build_full_presentation()
    short = build_short_presentation()
    print(f"Wrote full deck  ({full})")
    print(f"Wrote short deck ({short})")
